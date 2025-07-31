from django.views.generic import ListView, CreateView, UpdateView, DeleteView
from django.urls import reverse_lazy
from django.contrib import messages
from django.shortcuts import redirect
from django.db.models import Q
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .models import Plan
from .forms import PlanForm

# Mixins for filtering and exporting plans
class PlanFilterMixin:
    def apply_filters(self, request, queryset):
        status = request.GET.get('status')
        plan_type = request.GET.get('type')
        search = request.GET.get('search')

        if status in ['active', 'inactive']:
            queryset = queryset.filter(status=status)

        if plan_type in ['base', 'add_on']:
            queryset = queryset.filter(plan_type=plan_type)

        if search:
            queryset = queryset.filter(name__icontains=search)

        return queryset, status, plan_type, search

    def export_as_excel(self, request, queryset):
        if request.GET.get('export') == 'xlsx':
            wb = Workbook()
            ws = wb.active
            ws.title = "Plans"

            headers = ["ID", "Name", "Type", "Price", "Duration", "Status"]
            for col_num, header in enumerate(headers, 1):
                col_letter = get_column_letter(col_num)
                ws[f"{col_letter}1"] = header

            for row_num, plan in enumerate(queryset, 2):
                data = [
                    plan.id,
                    plan.name,
                    plan.plan_type,
                    plan.price,
                    plan.duration_days,
                    plan.status,
                ]
                for col_num, value in enumerate(data, 1):
                    col_letter = get_column_letter(col_num)
                    ws[f"{col_letter}{row_num}"] = value

            response = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
            response['Content-Disposition'] = 'attachment; filename=\"plans.xlsx\"'
            wb.save(response)
            return response
        return None

# Views for Plan Management
class PlanListView(PlanFilterMixin, ListView):
    model = Plan
    template_name = 'plans/plan-management.html'
    context_object_name = 'plans'

    def get_queryset(self):
        queryset = Plan.objects.all().order_by('-id')
        filtered_qs, _, _, _ = self.apply_filters(self.request, queryset)
        return filtered_qs

    def get(self, request, *args, **kwargs):
        queryset = self.get_queryset()
        export_response = self.export_as_excel(request, queryset)
        if export_response:
            return export_response
        return super().get(request, *args, **kwargs)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['form'] = PlanForm()
        _, status, plan_type, search = self.apply_filters(self.request, self.get_queryset())
        context.update({
            'status': status or '',
            'type': plan_type or '',
            'search': search or '',
        })
        return context

# Views for creating, updating, and deleting plans
class PlanCreateView(CreateView):
    model = Plan
    form_class = PlanForm
    success_url = reverse_lazy('plan:plan_management')

    def form_valid(self, form):
        messages.success(self.request, "Plan created successfully.")
        return super().form_valid(form)


class PlanUpdateView(UpdateView):
    model = Plan
    form_class = PlanForm
    success_url = reverse_lazy('plan:plan_management')
    context_object_name = 'plan'


class PlanDeleteView(DeleteView):
    model = Plan
    success_url = reverse_lazy('plan:plan_management')
    context_object_name = 'plan'

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        self.object.status = 'inactive'
        self.object.save()
        return redirect(self.success_url)

